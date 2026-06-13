"""s0_register 集成测试(连真 PG;不可达 skip)。临时构造批次,测完按 FK 序清理。"""

import io
from pathlib import Path

import pytest
from docx import Document as Docx
from openpyxl import Workbook
from sqlalchemy import delete, select, text
from ulid import ULID

from pipeline.config import load_config
from pipeline.index.object_store import ObjectStore
from pipeline.index.pg_io import PgIO
from pipeline.index.pg_models import Document, DocVersion, ImportBatch, PipelineEvent
from pipeline.stage_base import StageContext
from pipeline.stages.s0_register import REQUIRED_COLUMNS, register_batch


@pytest.fixture
def pg():
    io_ = PgIO.from_config(load_config())
    try:
        with io_.session() as s:
            s.execute(text("select 1"))
    except Exception:
        pytest.skip("PG 不可达(demo up 未起)")
    return io_


@pytest.fixture
def reg(pg, tmp_path):
    ctx = StageContext(config=load_config(), object_store=ObjectStore(tmp_path / "obj"), db=pg)
    batches: list[str] = []
    yield ctx, tmp_path, batches
    with pg.session() as s:
        dvs = list(s.scalars(select(DocVersion).where(DocVersion.batch_id.in_(batches or [""]))))
        dvids = [dv.doc_version_id for dv in dvs]
        lids = {dv.logical_id for dv in dvs}
        if dvids:
            s.execute(delete(PipelineEvent).where(PipelineEvent.doc_version_id.in_(dvids)))
            s.execute(delete(DocVersion).where(DocVersion.doc_version_id.in_(dvids)))
        if lids:
            s.execute(delete(Document).where(Document.logical_id.in_(lids)))
        if batches:
            s.execute(delete(ImportBatch).where(ImportBatch.batch_id.in_(batches)))


def _docx(text_: str = "第一条 测试内容") -> bytes:
    buf = io.BytesIO()
    d = Docx()
    d.add_paragraph(text_)
    d.save(buf)
    return buf.getvalue()


def _pdf() -> bytes:
    return b"%PDF-1.4\n% minimal demo pdf\n%%EOF\n"


def _xlsx() -> bytes:
    buf = io.BytesIO()
    wb = Workbook()
    wb.active["A1"] = "x"
    wb.save(buf)
    return buf.getvalue()


def _bid() -> str:
    return "t_" + str(ULID())  # 完整 ULID:前缀是时间戳,截断会在同毫秒内相撞


def _make_batch(base: Path, bid: str, rows: list[dict], files: dict, *, drop_col=None):
    d = base / bid
    d.mkdir(parents=True, exist_ok=True)
    for fn, data in files.items():
        (d / fn).write_bytes(data)
    cols = [c for c in REQUIRED_COLUMNS if c != drop_col]
    wb = Workbook()
    ws = wb.active
    ws.append(cols)
    for r in rows:
        ws.append([r.get(c, "") for c in cols])
    mp = d / "manifest.xlsx"
    wb.save(str(mp))
    return d, mp


def _row(filename, **kw):
    base = {
        "filename": filename, "title": f"标题-{filename}", "doc_number": f"令-{filename}",
        "issuer": "CSRC", "perm_tag": "公开", "corpus_type": "P-EXT",
        "biz_domain": "DISCLOSURE", "issue_date": "2024-01-01", "supersedes": "",
    }
    base.update(kw)
    return base


def test_manifest_missing_column_rejects_batch(reg):
    ctx, base, _ = reg
    bid = _bid()
    d, mp = _make_batch(base, bid, [_row("a.docx")], {"a.docx": _docx()}, drop_col="biz_domain")
    rep = register_batch(ctx, bid, d, mp)
    assert rep.accepted is False and "biz_domain" in rep.reject_reason
    assert ctx.db.get(ImportBatch, bid) is None  # 整批拒收,未落库


def test_registers_docx_and_pdf(reg):
    ctx, base, batches = reg
    bid = _bid()
    batches.append(bid)
    rows = [_row("a.docx", corpus_type="P-INT", perm_tag="内部"), _row("b.pdf")]
    d, mp = _make_batch(base, bid, rows, {"a.docx": _docx(), "b.pdf": _pdf()})
    rep = register_batch(ctx, bid, d, mp)
    assert rep.accepted and rep.counts()["REGISTERED"] == 2
    for o in rep.outcomes:
        dv = ctx.db.get(DocVersion, o.doc_version_id)
        assert dv.pipeline_status == "REGISTERED"
        assert len(o.doc_version_id) == 26  # ULID
        assert ctx.object_store.exists(dv.raw_object_key)  # 原件写一次


def test_whitelist_quarantine(reg):
    ctx, base, batches = reg
    bid = _bid()
    batches.append(bid)
    d, mp = _make_batch(base, bid, [_row("c.xlsx", corpus_type="P-INT")], {"c.xlsx": _xlsx()})
    o = register_batch(ctx, bid, d, mp).outcomes[0]
    assert o.status == "QUARANTINED" and o.error_code == "E101-DEMO"


def test_perm_missing_quarantine(reg):
    ctx, base, batches = reg
    bid = _bid()
    batches.append(bid)
    d, mp = _make_batch(base, bid, [_row("a.docx", perm_tag="")], {"a.docx": _docx()})
    o = register_batch(ctx, bid, d, mp).outcomes[0]
    assert o.status == "QUARANTINED" and "密级缺失" in o.reason


def test_sha_dedup_second_is_duplicate(reg):
    ctx, base, batches = reg
    data = _docx("同一内容")
    bid1 = _bid()
    batches.append(bid1)
    d1, mp1 = _make_batch(base, bid1, [_row("a.docx")], {"a.docx": data})
    register_batch(ctx, bid1, d1, mp1)
    bid2 = _bid()
    batches.append(bid2)
    d2, mp2 = _make_batch(base, bid2, [_row("a.docx")], {"a.docx": data})
    o = register_batch(ctx, bid2, d2, mp2).outcomes[0]
    assert o.status == "DUPLICATE"


def test_supersede_inherits_logical(reg):
    ctx, base, batches = reg
    bid1 = _bid()
    batches.append(bid1)
    d1, mp1 = _make_batch(base, bid1, [_row("v1.docx", doc_number="令1")], {"v1.docx": _docx("v1")})
    out1 = register_batch(ctx, bid1, d1, mp1).outcomes[0]
    bid2 = _bid()
    batches.append(bid2)
    r2 = [_row("v2.docx", doc_number="令2", supersedes="v1.docx")]
    d2, mp2 = _make_batch(base, bid2, r2, {"v2.docx": _docx("v2")})
    out2 = register_batch(ctx, bid2, d2, mp2).outcomes[0]
    assert out2.logical_id == out1.logical_id  # 替代 → logical 继承
    dv2 = ctx.db.get(DocVersion, out2.doc_version_id)
    assert dv2.supersedes_version_id == out1.doc_version_id
    assert dv2.version_relation == "revise_replace"
